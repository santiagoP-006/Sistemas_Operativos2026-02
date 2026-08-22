"""
Simulacion de cruce de cuatro vias con algoritmo del banquero.
4 direcciones, 4 semaforos, hilos por direccion.
Genera archivo Excel con resultados completos.

Modelo del banquero adaptado a trafico:
- Recursos: slots en la interseccion (capacidad limitada).
- Procesos: las direcciones de trafico activas en cada fase.
- Max claim por proceso: capacidad // direcciones activas.
- Cada vehiculo solicita 1 slot, cruza, y lo libera.
- El banquero verifica estado seguro antes de cada asignacion.
- Si el estado resultante es inseguro, la solicitud se rechaza.
- Los hilos pueden tener multiples vehiculos cruzando simultaneamente.
"""

import threading
import time
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

# --- Configuracion del cruce ---
VEHICULOS = {
    "A->B": 50,   # Izquierda a derecha
    "B->A": 35,   # Derecha a izquierda
    "C->D": 20,   # Arriba a abajo
    "D->C": 100   # Abajo a arriba
}

TIEMPO_VERDE = 3.0        # Segundos en verde
TIEMPO_AMARILLO = 1.0     # Segundos en amarillo
TIEMPO_PASO_AUTO = 0.05   # Tiempo que tarda un vehiculo en cruzar
CAPACIDAD_CRUCE = 4       # Slots simultaneos en la interseccion

# Fases sin conflicto
FASES = [
    ("HORIZONTAL", ["A->B", "B->A"]),
    ("VERTICAL", ["C->D", "D->C"])
]

SENTIDOS = {
    "A->B": "Izquierda a Derecha",
    "B->A": "Derecha a Izquierda",
    "C->D": "Arriba a Abajo",
    "D->C": "Abajo a Arriba"
}

# --- Registros globales ---
tiempo_inicio = None
tiempo_total_simulacion = 0
registro_vehiculos = []
registro_fases = []
registro_semaforos = []
registro_banquero = []
lock_registro = threading.Lock()


def log(msg):
    """Imprime con flush inmediato."""
    print(msg, flush=True)


# =============================================================
# ALGORITMO DEL BANQUERO (adaptado a trafico)
# =============================================================
class AlgoritmoBanquero:
    """
    Banquero adaptado a interseccion.
    Recursos = slots en el cruce (CAPACIDAD_CRUCE).
    Procesos = direcciones de trafico activas en cada fase.
    Max claim por proceso = CAPACIDAD_CRUCE // num_direcciones_activas.
    La necesidad real = max_claim - asignado.
    El banquero verifica estado seguro antes de cada asignacion.
    """

    def __init__(self, recursos_totales):
        self.recursos_totales = recursos_totales
        self.disponible = recursos_totales
        self.procesos_activos = {}   # {dir: vehiculos pendientes al inicio de fase}
        self.max_claim = {}          # {dir: max slots simultaneos permitidos}
        self.asignado = {}           # {dir: slots actualmente ocupados}
        self.lock = threading.Lock()

    def iniciar_fase(self, direcciones, pendientes, max_per_dir):
        """Configura el banquero al inicio de una fase verde.
        Solo incluye direcciones con vehiculos pendientes."""
        with self.lock:
            self.procesos_activos = {}
            self.max_claim = {}
            self.asignado = {}
            for d in direcciones:
                if pendientes.get(d, 0) > 0:
                    self.procesos_activos[d] = pendientes[d]
                    self.max_claim[d] = max_per_dir
                    self.asignado[d] = 0
            self.disponible = self.recursos_totales

    def actualizar_max_claim(self, direcciones_activas, nuevo_max):
        """Recalcula max_claim cuando una direccion termina.
        Se llama cuando una direccion se queda sin vehiculos pendientes
        para que la otra pueda usar mas slots."""
        with self.lock:
            # Quitar direcciones que ya terminaron
            self.procesos_activos = {
                d: v for d, v in self.procesos_activos.items()
                if d in direcciones_activas
            }
            # Actualizar max_claim de las que siguen activas
            for d in direcciones_activas:
                if d in self.max_claim:
                    self.max_claim[d] = nuevo_max

    def estado_seguro(self):
        """Verifica si el estado actual es seguro.
        Una direccion puede 'terminar' si su necesidad <= trabajo.
        Al terminar, libera todos sus slots asignados."""
        trabajo = self.disponible
        finalizado = {d: False for d in self.procesos_activos}
        necesidad = {}

        for d in self.procesos_activos:
            necesidad[d] = self.max_claim[d] - self.asignado[d]
            if necesidad[d] <= 0:
                finalizado[d] = True

        secuencia = []
        progreso = True
        while progreso:
            progreso = False
            for d in self.procesos_activos:
                if not finalizado[d] and necesidad[d] <= trabajo:
                    trabajo += self.asignado[d]
                    finalizado[d] = True
                    secuencia.append(d)
                    progreso = True

        return all(finalizado.values()), secuencia

    def solicitar(self, direccion):
        """Solicita 1 slot en la interseccion. True si concedido."""
        with self.lock:
            if direccion not in self.procesos_activos:
                return False
            if self.disponible < 1:
                self._registrar("RECHAZADO_SIN_SLOTS", direccion, [])
                return False

            # Verificar que no excede max_claim
            if self.asignado.get(direccion, 0) >= self.max_claim.get(direccion, 1):
                self._registrar("RECHAZADO_LIMITE", direccion, [])
                return False

            # Asignacion tentativa
            self.asignado[direccion] += 1
            self.disponible -= 1

            seguro, secuencia = self.estado_seguro()

            if seguro:
                self._registrar("CONCEDIDO", direccion, secuencia)
                return True

            # Revertir
            self.asignado[direccion] -= 1
            self.disponible += 1
            self._registrar("RECHAZADO_INSEGURO", direccion, [])
            return False

    def liberar(self, direccion):
        """Libera 1 slot tras cruzar."""
        with self.lock:
            if direccion in self.asignado and self.asignado[direccion] > 0:
                self.asignado[direccion] -= 1
                self.disponible += 1

    def _registrar(self, accion, direccion, secuencia):
        """Registra decision del banquero."""
        t = round(time.time() - tiempo_inicio, 3)
        with lock_registro:
            registro_banquero.append({
                "tiempo": t,
                "proceso": direccion,
                "accion": accion,
                "disponible_post": self.disponible,
                "asignado": dict(self.asignado),
                "max_claim": dict(self.max_claim),
                "secuencia": " -> ".join(secuencia) if secuencia else "N/A"
            })


# =============================================================
# SEMAFORO
# =============================================================
class Semaforo:
    """Semaforo con estados: VERDE, AMARILLO, ROJO."""

    def __init__(self, direccion):
        self.direccion = direccion
        self.estado = "ROJO"
        self.lock = threading.Lock()

    def cambiar(self, nuevo):
        """Cambia estado y registra la transicion solo si hay cambio real.
        Evita registrar transiciones innecesarias como ROJO->ROJO."""
        with self.lock:
            if self.estado == nuevo:
                return  # Sin cambio, no registrar
            anterior = self.estado
            self.estado = nuevo
        t = round(time.time() - tiempo_inicio, 3)
        with lock_registro:
            registro_semaforos.append({
                "tiempo": t,
                "direccion": self.direccion,
                "anterior": anterior,
                "nuevo": nuevo
            })

# =============================================================
# CONTROLADOR DEL CRUCE
# =============================================================
class ControladorCruce:
    """Coordina semaforos, hilos y el banquero.

    Cada hilo de direccion puede tener multiples vehiculos cruzando
    simultaneamente, limitados por el max_claim del banquero.
    El banquero ajusta dinamicamente el max_claim cuando una direccion
    termina antes que la otra.
    """

    def __init__(self):
        self.semaforos = {d: Semaforo(d) for d in VEHICULOS}
        self.pendientes = {d: VEHICULOS[d] for d in VEHICULOS}
        self.pasados = {d: 0 for d in VEHICULOS}
        self.en_cruce_total = 0
        self.lock_datos = threading.Lock()
        self.verde_activo = {d: threading.Event() for d in VEHICULOS}
        self.banquero = AlgoritmoBanquero(CAPACIDAD_CRUCE)
        self.ejecutando = True

    def todos_pasaron(self):
        """Verifica si todos los vehiculos terminaron de cruzar."""
        with self.lock_datos:
            return (all(self.pendientes[d] == 0 for d in VEHICULOS)
                    and self.en_cruce_total == 0)

    def pendientes_en(self, direccion):
        """Pendientes de una direccion."""
        with self.lock_datos:
            return self.pendientes[direccion]

    # --- Hilo por direccion ---
    def hilo_direccion(self, direccion):
        """Procesa vehiculos de una direccion cuando tiene verde.
        Soporta multiples vehiculos cruzando simultaneamente (pipeline)."""
        en_cruce = []  # Vehiculos actualmente en la interseccion

        while self.ejecutando:
            # Esperar senal de verde
            self.verde_activo[direccion].wait(timeout=0.5)
            if not self.ejecutando:
                break
            if not self.verde_activo[direccion].is_set():
                continue

            # --- Fase verde activa ---
            while self.verde_activo[direccion].is_set() and self.ejecutando:
                now = time.time()

                # 1. Liberar vehiculos que terminaron de cruzar
                completados = []
                restantes = []
                for vehiculo in en_cruce:
                    destino = (completados
                               if now - vehiculo["inicio"] >= TIEMPO_PASO_AUTO
                               else restantes)
                    destino.append(vehiculo)
                en_cruce = restantes
                for vehiculo in completados:
                    self.banquero.liberar(direccion)
                    salida = time.time()
                    vehiculo["registro"]["tiempo_salida"] = round(
                        salida - tiempo_inicio, 3)
                    vehiculo["registro"]["duracion"] = round(
                        salida - vehiculo["inicio"], 3)
                    with self.lock_datos:
                        self.en_cruce_total -= 1

                # 2. Solicitar nuevos vehiculos (el banquero es el unico limitante)
                solicitud_rechazada = False
                while self.verde_activo[direccion].is_set() and self.pendientes_en(direccion) > 0:
                    if self.banquero.solicitar(direccion):
                        with self.lock_datos:
                            if self.pendientes[direccion] > 0:
                                self.pendientes[direccion] -= 1
                                self.pasados[direccion] += 1
                                self.en_cruce_total += 1
                                num = self.pasados[direccion]
                            else:
                                self.banquero.liberar(direccion)
                                break
                        entrada = time.time()
                        t = round(entrada - tiempo_inicio, 3)
                        registro = {
                            "id": f"{direccion}_V{num}",
                            "direccion": direccion,
                            "sentido": SENTIDOS[direccion],
                            "numero": num,
                            "tiempo_entrada": t,
                            "tiempo_salida": None,
                            "duracion": None,
                            "semaforo": "VERDE"
                        }
                        en_cruce.append({"inicio": entrada, "registro": registro})
                        with lock_registro:
                            registro_vehiculos.append(registro)
                    else:
                        solicitud_rechazada = True
                        break  # Banquero rechazo, esperar al proximo ciclo

                # 3. Salir si no hay mas vehiculos ni en cruce
                if self.pendientes_en(direccion) <= 0 and len(en_cruce) == 0:
                    break

                if en_cruce:
                    espera = min(vehiculo["inicio"] + TIEMPO_PASO_AUTO - now
                                 for vehiculo in en_cruce)
                    time.sleep(max(0.001, min(0.02, espera)))
                elif solicitud_rechazada:
                    time.sleep(0.005)

            # --- Drenar vehiculos en cruce tras fin del verde ---
            while en_cruce and self.ejecutando:
                now = time.time()
                completados = []
                restantes = []
                for vehiculo in en_cruce:
                    destino = (completados
                               if now - vehiculo["inicio"] >= TIEMPO_PASO_AUTO
                               else restantes)
                    destino.append(vehiculo)
                en_cruce = restantes
                for vehiculo in completados:
                    self.banquero.liberar(direccion)
                    salida = time.time()
                    vehiculo["registro"]["tiempo_salida"] = round(
                        salida - tiempo_inicio, 3)
                    vehiculo["registro"]["duracion"] = round(
                        salida - vehiculo["inicio"], 3)
                    with self.lock_datos:
                        self.en_cruce_total -= 1
                if en_cruce:
                    espera = min(vehiculo["inicio"] + TIEMPO_PASO_AUTO - now
                                 for vehiculo in en_cruce)
                    time.sleep(max(0.001, min(0.02, espera)))

        log(f"  [Hilo {direccion}] Finalizado. Pasaron {self.pasados[direccion]} vehiculos.")

    # --- Hilo controlador de fases ---
    def hilo_controlador(self):
        """Alterna fases del semaforo hasta que todos crucen."""
        ciclo = 0

        while not self.todos_pasaron() and self.ejecutando:
            ciclo += 1
            log(f"\n--- Ciclo {ciclo} ---")

            for nombre_fase, dirs_verde in FASES:
                if self.todos_pasaron():
                    break

                # Calcular pendientes y direcciones activas
                with self.lock_datos:
                    pend_fase = {d: self.pendientes[d] for d in dirs_verde}
                dirs_activas = [d for d in dirs_verde if pend_fase[d] > 0]
                if not dirs_activas:
                    continue

                dirs_rojo = [d for d in VEHICULOS if d not in dirs_verde]
                fase_id = f"C{ciclo}_{nombre_fase[0]}"
                t_inicio = round(time.time() - tiempo_inicio, 3)

                # Calcular max concurrente por direccion activa
                n_active = len(dirs_activas)
                max_per = max(1, CAPACIDAD_CRUCE // n_active)

                # Configurar banquero para esta fase
                self.banquero.iniciar_fase(dirs_activas, pend_fase, max_per)

                log(f"  {fase_id}: {n_active} dirs activas, max_claim={max_per}")

                # --- VERDE ---
                for d in dirs_rojo:
                    self.semaforos[d].cambiar("ROJO")
                    self.verde_activo[d].clear()
                for d in dirs_verde:
                    self.semaforos[d].cambiar("VERDE")
                    self.verde_activo[d].set()
                inicio_verde = time.time()

                # Mantener verde (sale antes si no hay pendientes)
                fin_verde = time.time() + TIEMPO_VERDE
                while time.time() < fin_verde:
                    with self.lock_datos:
                        if all(self.pendientes[d] == 0 for d in dirs_verde):
                            break
                        # Detectar si una direccion termino para ajustar max_claim
                        active_now = [d for d in dirs_verde
                                      if self.pendientes[d] > 0]
                    if len(active_now) != n_active:
                        n_active = len(active_now)
                        new_max = max(1, CAPACIDAD_CRUCE // n_active)
                        self.banquero.actualizar_max_claim(active_now, new_max)
                        log(f"  Ajuste: {n_active} dirs activas, max_claim={new_max}")
                    time.sleep(0.02)

                # --- AMARILLO ---
                fin_verde_real = time.time()
                for d in dirs_verde:
                    self.verde_activo[d].clear()
                    self.semaforos[d].cambiar("AMARILLO")
                time.sleep(TIEMPO_AMARILLO)

                # --- ROJO ---
                for d in dirs_verde:
                    self.semaforos[d].cambiar("ROJO")

                t_fin = round(time.time() - tiempo_inicio, 3)

                # Contar vehiculos que cruzaron en esta fase
                pasados_fase = {d: 0 for d in dirs_verde}
                with lock_registro:
                    for v in registro_vehiculos:
                        if v["direccion"] in dirs_verde:
                            if t_inicio <= v["tiempo_entrada"] <= t_fin:
                                pasados_fase[v["direccion"]] += 1

                with lock_registro:
                    registro_fases.append({
                        "fase": fase_id,
                        "ciclo": ciclo,
                        "tipo": nombre_fase,
                        "direcciones": ", ".join(dirs_verde),
                        "t_inicio": t_inicio,
                        "t_fin": t_fin,
                        "duracion": round(t_fin - t_inicio, 3),
                        "inicio_verde": round(inicio_verde - tiempo_inicio, 3),
                        "fin_verde": round(fin_verde_real - tiempo_inicio, 3),
                        "duracion_verde": round(fin_verde_real - inicio_verde, 3),
                        "vehiculos": dict(pasados_fase),
                        "total": sum(pasados_fase.values())
                    })

                detalle = ", ".join(f"{d}:{pasados_fase[d]}" for d in dirs_verde)
                log(f"  {fase_id} [{nombre_fase}]: {detalle} = {sum(pasados_fase.values())} vehiculos")

                with self.lock_datos:
                    pend_str = ", ".join(f"{d}:{self.pendientes[d]}" for d in VEHICULOS)
                log(f"  Pendientes: {pend_str}")

                # Pausa entre fases (todo rojo)
                time.sleep(0.1)

    # --- Ejecucion principal ---
    def ejecutar(self):
        """Lanza todos los hilos y espera finalizacion."""
        log("Iniciando hilos de direccion...")
        hilos = []
        for d in VEHICULOS:
            t = threading.Thread(
                target=self.hilo_direccion, args=(d,),
                name=f"Dir-{d}", daemon=True
            )
            t.start()
            hilos.append(t)
            log(f"  Hilo {d} activo")

        log("\nControlador de semaforos iniciado.\n")
        ctrl = threading.Thread(target=self.hilo_controlador, name="Controlador")
        ctrl.start()
        ctrl.join()

        self.ejecutando = False
        for event in self.verde_activo.values():
            event.clear()
        for h in hilos:
            h.join()

        # Resultado final
        log("\n" + "=" * 55)
        log("RESULTADOS FINALES")
        log("=" * 55)
        total_p = 0
        total_e = 0
        for d in VEHICULOS:
            estado = "OK" if self.pasados[d] == VEHICULOS[d] else "FALTARON"
            log(f"  {d} ({SENTIDOS[d]}): {self.pasados[d]}/{VEHICULOS[d]} [{estado}]")
            total_p += self.pasados[d]
            total_e += VEHICULOS[d]
        log(f"  TOTAL: {total_p}/{total_e}")
        if total_p == total_e:
            log("  Estado: TODOS LOS VEHICULOS CRUZARON SIN CONFLICTOS")
        else:
            log("  Estado: SIMULACION INCOMPLETA")


# =============================================================
# GENERADOR DE EXCEL
# =============================================================
def generar_excel(controlador):
    """Genera Excel con 6 hojas de datos de la simulacion."""
    wb = Workbook()

    # Estilos
    borde = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    f_header = Font(bold=True, size=11)
    centrado = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def celda(ws, f, c, val, header=False, negrita=False):
        cell = ws.cell(row=f, column=c, value=val)
        cell.alignment = centrado
        cell.border = borde
        if header:
            cell.font = f_header
        else:
            font_args = {}
            if negrita:
                font_args["bold"] = True
            if font_args:
                cell.font = Font(**font_args)
        return cell

    def ancho_cols(ws, anchos):
        for letra, ancho in anchos.items():
            ws.column_dimensions[letra].width = ancho

    # ============================================
    # HOJA 1: RESUMEN GENERAL
    # ============================================
    ws1 = wb.active
    ws1.title = "Resumen"

    ws1.merge_cells('A1:F1')
    c1 = ws1['A1']
    c1.value = "SIMULACION CRUCE DE CUATRO VIAS - ALGORITMO DEL BANQUERO"
    c1.font = Font(bold=True, size=14)
    c1.alignment = centrado

    h1 = ["Direccion", "Sentido", "Total Vehiculos", "Vehiculos Pasados",
          "Pendientes", "Estado"]
    for i, h in enumerate(h1, 1):
        celda(ws1, 3, i, h, header=True)

    fila = 4
    for d in VEHICULOS:
        pend = VEHICULOS[d] - controlador.pasados[d]
        est = "COMPLETO" if pend == 0 else "PENDIENTE"
        celda(ws1, fila, 1, d)
        celda(ws1, fila, 2, SENTIDOS[d])
        celda(ws1, fila, 3, VEHICULOS[d])
        celda(ws1, fila, 4, controlador.pasados[d])
        celda(ws1, fila, 5, pend)
        celda(ws1, fila, 6, est, negrita=True)
        fila += 1

    celda(ws1, fila, 1, "TOTAL", negrita=True)
    celda(ws1, fila, 2, "")
    celda(ws1, fila, 3, sum(VEHICULOS.values()), negrita=True)
    celda(ws1, fila, 4, sum(controlador.pasados.values()), negrita=True)
    celda(ws1, fila, 5, sum(VEHICULOS.values()) - sum(controlador.pasados.values()), negrita=True)
    celda(ws1, fila, 6, "")

    fila += 2
    ws1.merge_cells(f'A{fila}:B{fila}')
    celda(ws1, fila, 1, "PARAMETROS DE SIMULACION", negrita=True)
    fila += 1
    params = [
        ("Duracion verde", f"{TIEMPO_VERDE} s"),
        ("Duracion amarillo", f"{TIEMPO_AMARILLO} s"),
        ("Tiempo de cruce por vehiculo", f"{TIEMPO_PASO_AUTO} s"),
        ("Capacidad de la interseccion", f"{CAPACIDAD_CRUCE} vehiculos"),
        ("Max concurrente por direccion (2 dirs)",
         f"{CAPACIDAD_CRUCE // 2} vehiculos"),
        ("Total de ciclos ejecutados",
         f"{len(set(f['ciclo'] for f in registro_fases))}"),
        ("Total de fases ejecutadas", f"{len(registro_fases)}"),
        ("Total de cambios de semaforo", f"{len(registro_semaforos)}"),
        ("Decisiones del banquero", f"{len(registro_banquero)}"),
        ("Tiempo total de la simulacion", f"{tiempo_total_simulacion} s"),
    ]
    for nombre, val in params:
        celda(ws1, fila, 1, nombre)
        celda(ws1, fila, 2, val)
        fila += 1

    fila += 1
    ws1.merge_cells(f'A{fila}:F{fila}')
    celda(ws1, fila, 1, "DESCRIPCION DEL CRUCE DE CUATRO VIAS", negrita=True)
    fila += 1
    desc = [
        "Cruce de cuatro vias con 2 carriles por sentido.",
        "Eje horizontal: Punto A (izquierda) <-> Punto B (derecha).",
        "Eje vertical: Punto C (arriba) <-> Punto D (abajo).",
        "A->B: 50 vehiculos, B->A: 35 vehiculos.",
        "C->D: 20 vehiculos, D->C: 100 vehiculos.",
        "Total: 205 vehiculos.",
        "4 semaforos (uno por direccion).",
        "Algoritmo del banquero evita deadlocks y controla capacidad.",
        "Multiples vehiculos por direccion cruzan simultaneamente.",
        "Max claim se ajusta dinamicamente si una direccion termina antes.",
        "Hilos independientes por cada direccion.",
    ]
    for linea in desc:
        celda(ws1, fila, 1, linea)
        ws1.merge_cells(f'A{fila}:F{fila}')
        ws1[f'A{fila}'].alignment = Alignment(horizontal="left", vertical="center")
        fila += 1

    ancho_cols(ws1, {'A': 38, 'B': 25, 'C': 18, 'D': 20, 'E': 14, 'F': 14})

    # ============================================
    # HOJA 2: FASES DEL SEMAFORO
    # ============================================
    ws2 = wb.create_sheet("Fases Semaforo")

    ws2.merge_cells('A1:L1')
    c2 = ws2['A1']
    c2.value = "DETALLE DE FASES DEL SEMAFORO POR CICLO"
    c2.font = Font(bold=True, size=13)
    c2.alignment = centrado

    h2 = ["Fase", "Ciclo", "Tipo", "Direcciones en Verde",
          "Inicio (s)", "Fin (s)", "Duracion (s)",
                    "Inicio Verde (s)", "Fin Verde (s)", "Duracion Verde (s)",
                    "Vehiculos por Direccion", "Total Vehiculos"]
    for i, h in enumerate(h2, 1):
        celda(ws2, 3, i, h, header=True)

    for idx, f in enumerate(registro_fases):
        r = idx + 4
        celda(ws2, r, 1, f["fase"])
        celda(ws2, r, 2, f["ciclo"])
        celda(ws2, r, 3, f["tipo"])
        celda(ws2, r, 4, f["direcciones"])
        celda(ws2, r, 5, f["t_inicio"])
        celda(ws2, r, 6, f["t_fin"])
        celda(ws2, r, 7, f["duracion"])
        celda(ws2, r, 8, f["inicio_verde"])
        celda(ws2, r, 9, f["fin_verde"])
        celda(ws2, r, 10, f["duracion_verde"])
        detalle_v = ", ".join(f"{k}: {v}" for k, v in f["vehiculos"].items())
        celda(ws2, r, 11, detalle_v)
        celda(ws2, r, 12, f["total"], negrita=True)

    r_total = len(registro_fases) + 4
    celda(ws2, r_total, 1, "TOTALES", negrita=True)
    for c in range(2, 12):
        celda(ws2, r_total, c, "")
    celda(ws2, r_total, 12, sum(f["total"] for f in registro_fases), negrita=True)

    ancho_cols(ws2, {'A': 10, 'B': 8, 'C': 14, 'D': 22, 'E': 12,
                     'F': 12, 'G': 14, 'H': 15, 'I': 15, 'J': 18,
                     'K': 30, 'L': 16})

    # ============================================
    # HOJA 3: VEHICULOS
    # ============================================
    ws3 = wb.create_sheet("Vehiculos")

    ws3.merge_cells('A1:H1')
    c3 = ws3['A1']
    c3.value = "REGISTRO DE CADA VEHICULO QUE CRUZO LA INTERSECCION"
    c3.font = Font(bold=True, size=13)
    c3.alignment = centrado

    h3 = ["ID Vehiculo", "Direccion", "Sentido", "Numero", "Entrada (s)", "Salida (s)", "Duracion (s)", "Semaforo"]
    for i, h in enumerate(h3, 1):
        celda(ws3, 3, i, h, header=True)

    for idx, v in enumerate(registro_vehiculos):
        r = idx + 4
        celda(ws3, r, 1, v["id"])
        celda(ws3, r, 2, v["direccion"])
        celda(ws3, r, 3, v["sentido"])
        celda(ws3, r, 4, v["numero"])
        celda(ws3, r, 5, v["tiempo_entrada"])
        celda(ws3, r, 6, v["tiempo_salida"])
        celda(ws3, r, 7, v["duracion"])
        celda(ws3, r, 8, v["semaforo"], negrita=True)

    r_sub = len(registro_vehiculos) + 5
    celda(ws3, r_sub, 1, "SUBTOTALES POR DIRECCION", negrita=True)
    r_sub += 1
    celda(ws3, r_sub, 1, "Direccion", header=True)
    celda(ws3, r_sub, 2, "Cantidad", header=True)
    celda(ws3, r_sub, 3, "Primer Cruce (s)", header=True)
    celda(ws3, r_sub, 4, "Ultimo Cruce (s)", header=True)
    r_sub += 1

    for d in VEHICULOS:
        tiempos = [v["tiempo_entrada"] for v in registro_vehiculos
                   if v["direccion"] == d]
        celda(ws3, r_sub, 1, d)
        celda(ws3, r_sub, 2, len(tiempos))
        celda(ws3, r_sub, 3, min(tiempos) if tiempos else "N/A")
        celda(ws3, r_sub, 4, max(tiempos) if tiempos else "N/A")
        r_sub += 1

    ancho_cols(ws3, {'A': 16, 'B': 12, 'C': 22, 'D': 10, 'E': 14,
                     'F': 14, 'G': 14, 'H': 12})

    # ============================================
    # HOJA 4: CAMBIOS DE SEMAFORO
    # ============================================
    ws4 = wb.create_sheet("Cambios Semaforo")

    ws4.merge_cells('A1:E1')
    c4 = ws4['A1']
    c4.value = "TRANSICIONES DE ESTADO DE LOS 4 SEMAFOROS"
    c4.font = Font(bold=True, size=13)
    c4.alignment = centrado

    h4 = ["Tiempo (s)", "Direccion", "Sentido", "Estado Anterior", "Estado Nuevo"]
    for i, h in enumerate(h4, 1):
        celda(ws4, 3, i, h, header=True)

    for idx, s in enumerate(registro_semaforos):
        r = idx + 4
        celda(ws4, r, 1, s["tiempo"])
        celda(ws4, r, 2, s["direccion"])
        celda(ws4, r, 3, SENTIDOS.get(s["direccion"], ""))
        celda(ws4, r, 4, s["anterior"], negrita=True)
        celda(ws4, r, 5, s["nuevo"], negrita=True)

    r_res = len(registro_semaforos) + 5
    celda(ws4, r_res, 1, "RESUMEN DE TRANSICIONES", negrita=True)
    r_res += 1
    celda(ws4, r_res, 1, "Direccion", header=True)
    celda(ws4, r_res, 2, "A Verde", header=True)
    celda(ws4, r_res, 3, "A Amarillo", header=True)
    celda(ws4, r_res, 4, "A Rojo", header=True)
    celda(ws4, r_res, 5, "Total", header=True)
    r_res += 1

    for d in VEHICULOS:
        cambios_d = [s for s in registro_semaforos if s["direccion"] == d]
        a_verde = sum(1 for s in cambios_d if s["nuevo"] == "VERDE")
        a_amar = sum(1 for s in cambios_d if s["nuevo"] == "AMARILLO")
        a_rojo = sum(1 for s in cambios_d if s["nuevo"] == "ROJO")
        celda(ws4, r_res, 1, d)
        celda(ws4, r_res, 2, a_verde)
        celda(ws4, r_res, 3, a_amar)
        celda(ws4, r_res, 4, a_rojo)
        celda(ws4, r_res, 5, len(cambios_d))
        r_res += 1

    ancho_cols(ws4, {'A': 14, 'B': 16, 'C': 22, 'D': 18, 'E': 18})

    # ============================================
    # HOJA 5: ALGORITMO DEL BANQUERO
    # ============================================
    ws5 = wb.create_sheet("Algoritmo Banquero")

    ws5.merge_cells('A1:H1')
    c5 = ws5['A1']
    c5.value = "ALGORITMO DEL BANQUERO - ASIGNACION DE RECURSOS"
    c5.font = Font(bold=True, size=13)
    c5.alignment = centrado

    # Explicacion
    celda(ws5, 3, 1, "MODELO DEL BANQUERO APLICADO A TRAFICO", negrita=True)
    ws5.merge_cells('A3:F3')
    explicaciones = [
        "Recursos: slots disponibles en la interseccion (capacidad limitada).",
        "Procesos: las direcciones de trafico activas en cada fase.",
        "Max claim por direccion: capacidad // direcciones activas.",
        "Necesidad real = max_claim - asignado (varia dinamicamente).",
        "El banquero verifica estado seguro antes de cada asignacion.",
        "Si el estado resultante es inseguro, la solicitud se rechaza.",
        "El max_claim se ajusta dinamicamente si una direccion termina antes.",
    ]
    for i, exp in enumerate(explicaciones):
        celda(ws5, 4 + i, 1, exp)
        ws5.merge_cells(f'A{4+i}:F{4+i}')
        ws5[f'A{4+i}'].alignment = Alignment(horizontal="left", vertical="center")

    # Estadisticas
    fila = 4 + len(explicaciones) + 1
    celda(ws5, fila, 1, "ESTADISTICAS", negrita=True)
    fila += 1
    concedidos = sum(1 for r in registro_banquero if r["accion"] == "CONCEDIDO")
    rechazados_slots = sum(1 for r in registro_banquero
                          if r["accion"] == "RECHAZADO_SIN_SLOTS")
    rechazados_inseg = sum(1 for r in registro_banquero
                           if r["accion"] == "RECHAZADO_INSEGURO")
    rechazados_limite = sum(1 for r in registro_banquero
                            if r["accion"] == "RECHAZADO_LIMITE")
    total_dec = (concedidos + rechazados_slots
                 + rechazados_inseg + rechazados_limite)

    stats = [
        ("Solicitudes concedidas", concedidos),
        ("Rechazadas (sin slots disponibles)", rechazados_slots),
        ("Rechazadas (limite por direccion)", rechazados_limite),
        ("Rechazadas (estado inseguro)", rechazados_inseg),
        ("Total de decisiones", total_dec),
    ]
    if total_dec > 0:
        stats.append(("Tasa de aprobacion",
                       f"{round(concedidos / total_dec * 100, 1)}%"))

    for nombre, val in stats:
        celda(ws5, fila, 1, nombre)
        celda(ws5, fila, 2, val)
        fila += 1

    # Concedidos por direccion
    fila += 1
    celda(ws5, fila, 1, "CONCEDIDOS POR DIRECCION", negrita=True)
    fila += 1
    celda(ws5, fila, 1, "Direccion", header=True)
    celda(ws5, fila, 2, "Concedidos", header=True)
    celda(ws5, fila, 3, "Rechazados", header=True)
    fila += 1
    for d in VEHICULOS:
        conc_d = sum(1 for r in registro_banquero
                     if r["proceso"] == d and r["accion"] == "CONCEDIDO")
        rech_d = sum(1 for r in registro_banquero
                     if r["proceso"] == d and r["accion"] != "CONCEDIDO")
        celda(ws5, fila, 1, d)
        celda(ws5, fila, 2, conc_d)
        celda(ws5, fila, 3, rech_d)
        fila += 1

    # Log de decisiones (muestra)
    fila += 1
    celda(ws5, fila, 1, "REGISTRO DE DECISIONES (muestra)", negrita=True)
    fila += 1
    h5b = ["Tiempo (s)", "Proceso", "Accion",
           "Disponible Post", "Secuencia Segura"]
    for i, h in enumerate(h5b, 1):
        celda(ws5, fila, i, h, header=True)
    fila += 1

    muestra = registro_banquero[:800]
    for reg in muestra:
        celda(ws5, fila, 1, reg["tiempo"])
        celda(ws5, fila, 2, reg["proceso"])
        celda(ws5, fila, 3, reg["accion"], negrita=True)
        celda(ws5, fila, 4, reg["disponible_post"])
        celda(ws5, fila, 5, reg["secuencia"])
        fila += 1

    ancho_cols(ws5, {'A': 28, 'B': 16, 'C': 24, 'D': 18, 'E': 40})

    # ============================================
    # HOJA 6: TIEMPOS DE DEMORA
    # ============================================
    ws6 = wb.create_sheet("Tiempos de Demora")

    ws6.merge_cells('A1:F1')
    c6 = ws6['A1']
    c6.value = "ANALISIS DE TIEMPOS DE DEMORA POR CAMBIO DE ESTADO"
    c6.font = Font(bold=True, size=13)
    c6.alignment = centrado

    celda(ws6, 3, 1, "DEMORA POR DIRECCION", negrita=True)
    h6 = ["Direccion", "Sentido", "Tiempo en Verde (s)",
          "Tiempo en Amarillo (s)", "Tiempo en Rojo (s)",
          "Vehiculos/Segundo"]
    for i, h in enumerate(h6, 1):
        celda(ws6, 4, i, h, header=True)

    fila = 5
    for d in VEHICULOS:
        cambios = [s for s in registro_semaforos if s["direccion"] == d]
        t_verde = 0.0
        t_amar = 0.0
        t_rojo = 0.0

        ultimo_t = 0
        if registro_vehiculos:
            ultimo_t = max(v["tiempo_salida"] for v in registro_vehiculos)

        for i in range(len(cambios)):
            inicio = cambios[i]["tiempo"]
            fin = (cambios[i + 1]["tiempo"]
                   if i + 1 < len(cambios) else ultimo_t)
            dur = max(0, fin - inicio)
            if cambios[i]["nuevo"] == "VERDE":
                t_verde += dur
            elif cambios[i]["nuevo"] == "AMARILLO":
                t_amar += dur
            elif cambios[i]["nuevo"] == "ROJO":
                t_rojo += dur

        vel = (round(controlador.pasados[d] / t_verde, 2)
               if t_verde > 0 else 0)
        celda(ws6, fila, 1, d)
        celda(ws6, fila, 2, SENTIDOS[d])
        celda(ws6, fila, 3, round(t_verde, 3))
        celda(ws6, fila, 4, round(t_amar, 3))
        celda(ws6, fila, 5, round(t_rojo, 3))
        celda(ws6, fila, 6, vel)
        fila += 1

    # Tiempo promedio por vehiculo
    fila += 2
    celda(ws6, fila, 1, "TIEMPO PROMEDIO DE CRUCE POR VEHICULO", negrita=True)
    fila += 1
    h6b = ["Direccion", "Primer Vehiculo (s)", "Ultimo Vehiculo (s)",
           "Tiempo Total Carril (s)", "Tiempo Acumulado Coches (s)",
           "Promedio entre Vehiculos (s)"]
    for i, h in enumerate(h6b, 1):
        celda(ws6, fila, i, h, header=True)
    fila += 1

    for d in VEHICULOS:
        tiempos = sorted([v["tiempo_entrada"] for v in registro_vehiculos
                          if v["direccion"] == d])
        salidas = [v["tiempo_salida"] for v in registro_vehiculos
               if v["direccion"] == d]
        if len(tiempos) >= 2:
            prim = tiempos[0]
            ult = tiempos[-1]
            dur = round(max(salidas) - prim, 3)
            acumulado = round(sum(v["duracion"] for v in registro_vehiculos
                                  if v["direccion"] == d), 3)
            prom = (round((ult - prim) / (len(tiempos) - 1), 4)
                    if len(tiempos) > 1 else 0)
        elif len(tiempos) == 1:
            prim = ult = tiempos[0]
            dur = acumulado = prom = 0
        else:
            prim = ult = dur = acumulado = prom = 0

        celda(ws6, fila, 1, d)
        celda(ws6, fila, 2, prim)
        celda(ws6, fila, 3, ult)
        celda(ws6, fila, 4, dur)
        celda(ws6, fila, 5, acumulado)
        celda(ws6, fila, 6, prom)
        fila += 1

    # Demora al cambiar estados
    fila += 2
    celda(ws6, fila, 1, "DEMORA AL CAMBIAR ESTADOS DEL SEMAFORO", negrita=True)
    fila += 1
    h6c = ["Transicion", "Direccion", "Tiempo (s)", "Efecto en Vehiculos"]
    for i, h in enumerate(h6c, 1):
        celda(ws6, fila, i, h, header=True)
    fila += 1

    for s in registro_semaforos:
        trans = f"{s['anterior']} -> {s['nuevo']}"
        efecto = ""
        if s["nuevo"] == "ROJO":
            efecto = "Vehiculos detenidos hasta proximo verde"
        elif s["nuevo"] == "AMARILLO":
            efecto = "Vehiculos dejan de entrar al cruce"
        elif s["nuevo"] == "VERDE":
            efecto = "Vehiculos comienzan a cruzar"

        celda(ws6, fila, 1, trans)
        celda(ws6, fila, 2, s["direccion"])
        celda(ws6, fila, 3, s["tiempo"])
        celda(ws6, fila, 4, efecto)
        fila += 1

    ancho_cols(ws6, {'A': 24, 'B': 22, 'C': 22, 'D': 20, 'E': 28, 'F': 18})

    # Guardar
    archivo = Path(__file__).resolve().with_name("simulacion_cruce_T.xlsx")
    try:
        wb.save(archivo)
    except PermissionError:
        marca_tiempo = datetime.now().strftime("%Y%m%d_%H%M%S")
        archivo = archivo.with_name(f"simulacion_cruce_T_{marca_tiempo}.xlsx")
        wb.save(archivo)
    log(f"\nArchivo Excel generado correctamente en:\n{archivo}")
    return archivo


# =============================================================
# MAIN
# =============================================================
def main():
    global tiempo_inicio, tiempo_total_simulacion

    registro_vehiculos.clear()
    registro_fases.clear()
    registro_semaforos.clear()
    registro_banquero.clear()

    log("=" * 55)
    log("SIMULACION CRUCE DE CUATRO VIAS - ALGORITMO DEL BANQUERO")
    log("=" * 55)
    log(f"\nVehiculos por direccion:")
    for d, c in VEHICULOS.items():
        log(f"  {d} ({SENTIDOS[d]}): {c}")
    log(f"  Total: {sum(VEHICULOS.values())}")
    log(f"\nSemaforos: 4 (uno por direccion)")
    log(f"Verde: {TIEMPO_VERDE}s | Amarillo: {TIEMPO_AMARILLO}s")
    log(f"Capacidad interseccion: {CAPACIDAD_CRUCE} vehiculos")
    log(f"Max concurrente por dir (2 dirs): {CAPACIDAD_CRUCE // 2}")
    log("")

    controlador = ControladorCruce()
    tiempo_inicio = time.time()

    controlador.ejecutar()

    tiempo_total_simulacion = round(time.time() - tiempo_inicio, 2)
    log(f"\nTiempo total de simulacion: {tiempo_total_simulacion} segundos")
    log(f"Eventos: {len(registro_vehiculos)} vehiculos, "
        f"{len(registro_semaforos)} cambios de semaforo, "
        f"{len(registro_banquero)} decisiones del banquero")

    generar_excel(controlador)
    log("Simulacion finalizada.")


if __name__ == "__main__":
    main()
